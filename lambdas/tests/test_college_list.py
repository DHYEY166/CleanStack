"""Regression tests from a real run on a small college list (xlsx) where every suggested
rule was approved and the output:

1. lost the 'Brown' row, whose three numeric cells held the placeholder ".";
2. carried orig_<col> text copies of three numeric columns although no value changed;
3. lost the 'COLLEGEPOND REFERENCE' section label (merged A2:F2, right under the header).

The fixture has the same shape: a header, a merged section label in row 2, blank spacer
rows, two more section labels, '.' placeholders, padded text and integer cells.
"""
import io

import openpyxl
import pandas as pd
import pytest

from conftest import rule

U = "UNIVERESITY"
WR = "WORLD RANKING(FOR DATA SCIENCE)"
NAR = "NORTH AMERICA RANKING(FOR DATA SCIENCE)"
AR = "Acceptance Rate(%)"
HEADER = [U, WR, NAR, "LOCATION", "TYPE", "CHANCES", AR]
ROWS = [
    ["COLLEGEPOND REFERENCE"],                                                  # merged A2:F2
    [],
    ["University of California, San Diago", 12, 10, "San Diego, California", "State University", "Super Ambitious", 34],
    [" Columbia University", 24, 17, "New York", "Ivy League", "Super Ambitious", 4],
    [],
    ["Duke", 64, 42, "North Carolina", "Private University", "Ambitious", 6],
    ["Brown", ".", ".", "Providence, Rhode Island", "Ivy League", "Ambitious", "."],
    [],
    ["MY PRIORITY"],                                                            # merged
    [],
    ["University of Washington", 5, 3, "Seattle, Washington", "State University", None, 53],
    ["University of Illinois at Urbana - Champaign", 16, 12, "Champaign, Illinois", "State University", None, None],
    ["LESS CHANCE BUT MY PRIORITY"],                                            # merged
    [],
    ["Harvard", 3, 2, "Cambridge, Massachusetts", "Ivy League", None, 4],
]
LABELS = ["COLLEGEPOND REFERENCE", "MY PRIORITY", "LESS CHANCE BUT MY PRIORITY"]
BLANK_ROWS = sum(1 for r in ROWS if not r)


def college_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(HEADER)
    for r in ROWS:
        ws.append(r)
    ws.merge_cells("A2:F2")
    ws.merge_cells("A10:F10")
    ws.merge_cells("A14:G14")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# What the AI suggested for the real file, minus any row filter: no rule in this list
# targets rows by value, so every non-blank row must survive.
SUGGESTED = [
    rule("trim_whitespace", rule_id="trim"),
    rule("deduplicate", rule_id="dedup"),
    rule("drop_nulls", U, rule_id="drop-blank"),
    rule("type_cast", AR, rule_id="cast-ar", target_type="float"),
    rule("type_cast", WR, rule_id="cast-wr", target_type="int"),
    rule("type_cast", NAR, rule_id="cast-nar", target_type="int"),
    rule("column_header_normalize", rule_id="headers"),
]


@pytest.fixture
def raw(executor):
    return executor.load_raw_dataframe(college_xlsx(), "xlsx")


def run(executor, df, rules):
    results = []
    return executor.apply_transforms(df.copy(), rules, results), results


# ---------------------------------------------------------------- 1. Brown
def test_placeholder_cells_become_null_and_the_row_is_kept(executor, raw):
    out, res = run(executor, raw, [rule("type_cast", c, target_type="int") for c in (WR, NAR, AR)])
    assert len(out) == len(raw)
    brown = out[out[U] == "Brown"].iloc[0]
    assert all(pd.isna(brown[c]) for c in (WR, NAR, AR))
    assert brown["LOCATION"] == "Providence, Rhode Island"
    assert all(r["applied"] and r["rows_removed"] == 0 for r in res)


def test_suggested_rules_keep_brown_and_every_label_row(executor, raw):
    out, res = run(executor, raw, SUGGESTED)
    names = out["univeresity"].tolist()
    assert "Brown" in names
    for label in LABELS:
        assert label in names
    assert len(out) == len(raw) - BLANK_ROWS
    removed = {r["id"]: r["rows_removed"] for r in res}
    # Only the blank spacer rows go: the first by drop_nulls, the repeats by deduplicate.
    assert removed["dedup"] + removed["drop-blank"] == BLANK_ROWS
    assert all(v == 0 for k, v in removed.items() if k not in ("dedup", "drop-blank"))


def test_an_explicit_row_filter_still_works_and_reports_what_it_removed(executor, raw):
    out, res = run(executor, raw, [rule("filter", WR, rule_id="f", operator="neq", value=".")])
    assert "Brown" not in out[U].tolist()
    assert res == [{"id": "f", "rule_type": "filter", "column_name": WR,
                    "applied": True, "reason": None, "rows_removed": 1, "bad_cell_rows": 1}]


def test_auto_mode_guard_keeps_brown(executor, raw):
    # Nobody approved the rule in auto mode: the bad-cell guard skips it and keeps the row.
    results = []
    out = executor.apply_transforms(raw.copy(), [rule("filter", WR, operator="neq", value=".")],
                                    results, skip_bad_cell_rules=True)
    assert "Brown" in out[U].tolist()
    assert results[0]["applied"] is False and results[0]["bad_cell_rows"] == 1


def test_rows_removed_is_zero_for_a_rule_that_was_not_applied(executor, raw):
    # drop_nulls on CHANCES would remove most rows: the row-loss guard skips it.
    out, res = run(executor, raw, [rule("drop_nulls", "CHANCES")])
    assert len(out) == len(raw)
    assert res[0]["applied"] is False and res[0]["rows_removed"] == 0


# ---------------------------------------------------------------- 2. orig_<col> copies
def test_trim_whitespace_leaves_numbers_numeric(executor, raw):
    out, _ = run(executor, raw, [rule("trim_whitespace")])
    assert out.loc[out[U] == "Duke", WR].iloc[0] == 64
    assert out.loc[out[U] == "Columbia University", U].iloc[0] == "Columbia University"


def test_lossless_cast_adds_no_sidecar(executor):
    df = pd.DataFrame({"rank": [12, 24, None], "rate": ["34", "4", None]})
    out, _ = run(executor, df, [rule("type_cast", "rank", target_type="int"),
                                rule("type_cast", "rate", target_type="float")])
    assert list(out.columns) == ["rank", "rate"]
    assert out["rate"].tolist()[:2] == [34.0, 4.0]


@pytest.mark.parametrize("original", [".", "$1,200", "007", " 12", "1e3"])
def test_lossy_cast_keeps_the_original_in_a_sidecar(executor, original):
    df = pd.DataFrame({"v": ["5", original]})
    out, _ = run(executor, df, [rule("type_cast", "v", target_type="float")])
    assert out["__orig_v"].tolist() == ["5", original]


def test_noop_fill_and_outlier_cap_add_no_sidecar(executor):
    df = pd.DataFrame({"x": [1.0, 2.0, 3.0, 2.5]})
    out, _ = run(executor, df, [rule("fill_nulls", "x", strategy="mean"),
                                rule("outlier_cap", "x", min_val=0, max_val=10),
                                rule("ffill", "x")])
    assert list(out.columns) == ["x"]


def test_header_normalize_keeps_sidecars_out_of_the_deliverable(executor, raw):
    # Brown's "." becoming null is a real loss, so these sidecars are kept, in audit.csv only.
    out, _ = run(executor, raw, SUGGESTED)
    assert "__orig_acceptance_rate" in out.columns
    assert not [c for c in out.columns if str(c).startswith("orig_")]
    deliverable, audit = executor.split_deliverable(out)
    assert list(deliverable.columns) == ["univeresity", "world_ranking_for_data_science",
                                         "north_america_ranking_for_data_science", "location",
                                         "type", "chances", "acceptance_rate"]
    assert audit is not None
    brown = audit[audit["univeresity"] == "Brown"].iloc[0]
    assert brown["__orig_world_ranking_for_data_science"] == "."


def test_rename_moves_the_sidecar_with_its_column(executor):
    df = pd.DataFrame({"Amount": ["$5", "7"]})
    out, _ = run(executor, df, [rule("type_cast", "Amount", target_type="float"),
                                rule("rename", "Amount", new_name="amount_usd")])
    assert list(out.columns) == ["amount_usd", "__orig_amount_usd"]


# ---------------------------------------------------------------- 3. COLLEGEPOND REFERENCE
def test_section_label_under_the_header_is_read_as_a_data_row(executor, raw):
    assert list(raw.columns) == HEADER
    assert raw.iloc[0][U] == "COLLEGEPOND REFERENCE"
    assert raw[U].tolist().count("COLLEGEPOND REFERENCE") == 1


def test_profiler_also_reads_the_label_as_a_data_row(profiler):
    df = profiler.load_dataframe(college_xlsx(), "xlsx")
    assert list(df.columns) == HEADER
    assert df.iloc[0][U] == "COLLEGEPOND REFERENCE"


def test_comparison_filter_keeps_rows_whose_value_is_missing_or_unparseable(executor, raw):
    # Before: num > 0 was False for "." and null, so Brown, the labels and the spacer rows
    # were all removed by a rule meant to drop non-positive rankings.
    out, res = run(executor, raw, [rule("filter", WR, rule_id="f", operator="gt", value=0)])
    assert len(out) == len(raw) and res[0]["rows_removed"] == 0
    df = pd.DataFrame({"r": [5, -1, None, ".", 7, 8]})
    out, res = run(executor, df, [rule("filter", "r", operator="gt", value=0)])
    assert -1 not in out["r"].tolist() and len(out) == 5 and res[0]["rows_removed"] == 1


@pytest.mark.parametrize("target", ["int", "float"])
def test_type_cast_null_values_turns_numeric_codes_into_null_and_keeps_the_row(executor, target):
    df = pd.DataFrame({"score": [12, 99999, "99999", -9999.0, "N/A", 7], "id": range(6)})
    out, res = run(executor, df, [rule("type_cast", "score", target_type=target, null_values=[99999, -9999, "n/a"])])
    assert len(out) == 6 and res[0]["rows_removed"] == 0
    assert out["score"].isna().tolist() == [False, True, True, True, True, False]
    assert out["__orig_score"].tolist()[1:3] == [99999, "99999"]


def test_type_cast_without_null_values_is_unchanged(executor):
    df = pd.DataFrame({"score": [12, 99999]})
    out, _ = run(executor, df, [rule("type_cast", "score", target_type="int")])
    assert out["score"].tolist() == [12, 99999] and list(out.columns) == ["score"]
